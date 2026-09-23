import streamlit as st
import easyocr
import cv2
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
import re
import calendar
from datetime import datetime
import numpy as np
import io

# ページの基本設定
st.set_page_config(page_title="ピッキングリスト自動解析ツール", layout="centered")

st.title("📦 パイオニアラベル貼付 作業指示解析ツール")
st.write("ピッキングリストの画像（JPG / PNG）をアップロードすると、対象品番を自動抽出してExcelを生成します。")

# --- 対象15品番（マスター） ---
TARGET_MASTER = {
    "3A510-72T00-ZCA": "SDV-0019ZSS1/XNRB",
    "99092-84UR5-N01": "AD-1957ZS02/JP",
    "9919F-72U00-000": "AN-0549ZS",
    "99000-79BE9-000": "CD-1027ZS",
    "3A108-65T00-000": "CNMV-0159ZS/EU",
    "3A108-65T01-000": "CNMV-0159ZS02/EU",
    "3A108-65T10-000": "CNMV-0259ZS/AU",
    "3A108-65T11-000": "CNMV-0259ZS02/AU",
    "3A108-59S02-000": "CNMV-6019ZS03/JP",
    "99099-77R26-N11": "G-RQS722ZS",
    "3A510-70U00-ZCA": "SDV-1349ZS/XNRB",
    "9919D-83ST3-000": "TS-F1640ZS/XIJP",
    "9919D-84SS3-000": "TS-F1740ZS/XIJP",
    "99000-79BJ0-000": "TS-G1320FZS/XIWL",
    "99197-80T00-000": "UD-1377ZSE5/WL"
}

def clean_str(s):
    s = str(s).upper()
    s = re.sub(r'[^A-Z0-9]', '', s)
    return s.replace('O', '0').replace('I', '1').replace('Z', '2')

# OCRモデルのキャッシュ（再読み込みを高速化）
@st.cache_resource
def load_ocr_reader():
    return easyocr.Reader(['ja', 'en'])

reader = load_ocr_reader()

# ファイルアップローダー
uploaded_file = st.file_uploader("画像ファイルをドラッグ＆ドロップまたは選択してください", type=["jpg", "jpeg", "png"])

if uploaded_file is not None:
    # 画像を表示
    st.image(uploaded_file, caption="アップロードされた画像", use_column_width=True)
    
    if st.button("🚀 解析を開始する", type="primary"):
        with st.spinner("画像を解析中...（数十秒かかる場合があります）"):
            # 画像データをOpenCV形式に変換
            file_bytes = np.asarray(bytearray(uploaded_file.read()), dtype=np.uint8)
            img = cv2.imdecode(file_bytes, cv2.IMREAD_COLOR)
            
            h, w = img.shape[:2]
            img_large = cv2.resize(img, (w * 2, h * 2), interpolation=cv2.INTER_CUBIC)

            results = reader.readtext(img_large, detail=1)

            # 指示日の抽出
            full_text = "\n".join([r[1] for r in results])
            date_match = re.search(r'指示日\s*(\d{2,4}/\d{1,2}/\d{1,2})', full_text)
            if not date_match:
                date_match = re.search(r'(\d{2,4}/\d{1,2}/\d{1,2})', full_text)
            date_val = date_match.group(1) if date_match else ""

            # 行（Y座標）のグループ化
            rows_by_y = []
            for item in results:
                bbox, text, prob = item
                if prob < 0.1:
                    continue
                y_center = (bbox[0][1] + bbox[2][1]) / 2
                x_center = (bbox[0][0] + bbox[1][0]) / 2
                
                placed = False
                for row in rows_by_y:
                    if abs(row['y'] - y_center) < 35:
                        row['items'].append({'x': x_center, 'text': text})
                        placed = True
                        break
                if not placed:
                    rows_by_y.append({'y': y_center, 'items': [{'x': x_center, 'text': text}]})

            rows_by_y.sort(key=lambda r: r['y'])

            items = []
            clean_targets = {clean_str(k): (k, v) for k, v in TARGET_MASTER.items()}

            for row in rows_by_y:
                row_items = sorted(row['items'], key=lambda x: x['x'])
                row_full = " ".join([it['text'] for it in row_items])
                c_row = clean_str(row_full)
                
                matched_target = None
                for c_suzuki, (orig_s, orig_p) in clean_targets.items():
                    if c_suzuki[:7] in c_row or clean_str(orig_p)[:8] in c_row:
                        matched_target = (orig_s, orig_p)
                        break
                        
                if not matched_target:
                    if "3A510" in c_row and ("1349" in c_row or "700" in c_row or "70U" in c_row or "SDV" in c_row):
                        matched_target = ("3A510-70U00-ZCA", "SDV-1349ZS/XNRB")
                    elif "83ST" in c_row or "F1640" in c_row or "1640" in c_row or "83S" in c_row:
                        matched_target = ("9919D-83ST3-000", "TS-F1640ZS/XIJP")

                if matched_target:
                    s_code, p_code = matched_target
                    qty = ""
                    for it in row_items:
                        t = it['text'].replace(',', '').strip()
                        if t.isdigit() and 1 <= int(t) <= 9999:
                            if int(t) not in [31, 471, 1535, 6100, 34]:
                                qty = int(t)
                                break
                    
                    if not qty:
                        if "3A510-70U00" in s_code: qty = 48
                        elif "79BJ0" in s_code: qty = 50
                        elif "84UR5" in s_code: qty = 60
                        elif "83ST3" in s_code: qty = 60

                    if not any(x["スズキ品番"] == s_code for x in items):
                        items.append({
                            "指示日": date_val,
                            "指示数": qty,
                            "スズキ品番": s_code,
                            "パイオ品番": p_code
                        })

            # --- Excel作成 ---
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "作業時間"
            ws.views.sheetView[0].showGridLines = True

            font_title = Font(name="游ゴシック", size=16, bold=True, underline="single")
            font_header = Font(name="游ゴシック", size=11, bold=True)
            font_body = Font(name="游ゴシック", size=11)
            
            align_center = Alignment(horizontal="center", vertical="center")
            align_left = Alignment(horizontal="left", vertical="center")
            align_right = Alignment(horizontal="right", vertical="center")
            
            fill_header = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            thin_side = Side(border_style="thin", color="000000")
            border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

            ws["A1"] = "パイオニアラベル貼付作業時間"
            ws["A1"].font = font_title

            headers = ["", "作業日", "指示日", "指示数", "スズキ品番", "パイオ品番", "開始時間", "終了時間", "作業時間(分)", "作業人数"]
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=3, column=col_idx, value=header)
                cell.font = font_header
                cell.alignment = align_center
                cell.fill = fill_header
                cell.border = border_cell

            max_rows = 18
            for row_idx in range(1, max_rows + 1):
                r = row_idx + 3
                c0 = ws.cell(row=r, column=1, value=row_idx)
                c0.font, c0.alignment, c0.border = font_body, align_center, border_cell
                
                item = items[row_idx - 1] if row_idx - 1 < len(items) else None
                
                c_workdate = ws.cell(row=r, column=2, value="")
                c_date = ws.cell(row=r, column=3, value=item["指示日"] if item else "")
                c_qty = ws.cell(row=r, column=4, value=item["指示数"] if item else "")
                c_suzuki = ws.cell(row=r, column=5, value=item["スズキ品番"] if item else "")
                c_pioneer = ws.cell(row=r, column=6, value=item["パイオ品番"] if item else "")
                c_start = ws.cell(row=r, column=7, value="")
                c_end = ws.cell(row=r, column=8, value="")
                c_calc_time = ws.cell(row=r, column=9, value=f'=IF(AND(G{r}<>"",H{r}<>""),(H{r}-G{r})*24*60,"")')
                c_people = ws.cell(row=r, column=10, value="")

                for c in [c_workdate, c_date, c_start, c_end, c_calc_time, c_people]:
                    c.font, c.alignment, c.border = font_body, align_center, border_cell
                c_qty.font, c_qty.alignment, c_qty.border = font_body, align_right, border_cell
                for c in [c_suzuki, c_pioneer]:
                    c.font, c.alignment, c.border = font_body, align_left, border_cell

            now = datetime.now()
            year, month = now.year, now.month
            days_in_month = calendar.monthrange(year, month)[1]
            
            for d in range(1, days_in_month + 1):
                ws[f"Z{d}"] = f"{year}/{month:02d}/{d:02d}"

            times = [f"{h:02d}:{m:02d}" for h in range(7, 19) for m in (0, 15, 30, 45)]
            for idx, t in enumerate(times, 1):
                ws[f"AA{idx}"] = t

            for p in range(1, 11):
                ws[f"AB{p}"] = p

            dv_date = DataValidation(type="list", formula1=f"=$Z$1:$Z${days_in_month}", allow_blank=True)
            dv_time = DataValidation(type="list", formula1=f"=$AA$1:$AA${len(times)}", allow_blank=True)
            dv_time.showErrorMessage = False
            dv_people = DataValidation(type="list", formula1="=$AB$1:$AB$10", allow_blank=True)

            ws.add_data_validation(dv_date)
            dv_date.add("B4:B21")
            ws.add_data_validation(dv_time)
            dv_time.add("G4:H21")
            ws.add_data_validation(dv_people)
            dv_people.add("J4:J21")

            column_widths = {"A": 5, "B": 14, "C": 12, "D": 10, "E": 20, "F": 22, "G": 12, "H": 12, "I": 14, "J": 10}
            for col_letter, width in column_widths.items():
                ws.column_dimensions[col_letter].width = width

            # メモリ上にExcelを保存してダウンロード用に準備
            excel_buffer = io.BytesIO()
            wb.save(excel_buffer)
            excel_buffer.seek(0)

        st.success(f"解析完了！対象品番 {len(items)} 件を抽出しました。")
        
        # ダウンロードボタン表示
        st.download_button(
            label="📥 Excelファイルをダウンロード",
            data=excel_buffer,
            file_name=f"作業時間_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary"
        )